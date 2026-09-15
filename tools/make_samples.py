#!/usr/bin/env python3
"""샘플 서류 22건 생성기 (2026-09-10 실측용 축소판, 리서치 `tax-office-document-labels` §4.2·§4.3 규격).

    python3 tools/make_samples.py generate   # samples/inbox/ 에 파일 생성 (+ hwpx 내용 JSON 2건)
    python3 tools/make_samples.py extract    # samples/text/*.txt + samples/labels.json

- 서류 종류 15종 실무 약칭(리서치 §4.2 enum). 인바운드 핵심 9종 × 2건 + 전자세금계산서 1건 + 사무소 산출물 3종 × 1건 = 22건.
- 포맷: 텍스트 추출 가능(xlsx·docx·텍스트 PDF·hwpx) 약 60% + OCR 필요(스캔 PDF·카톡 사진 JPG) 약 40%.
- 가상 고객사 5곳, 상호·대표자·사업자번호·주소·전화는 전부 가짜(실존 상호 중복은 오프라인이라 미확인).
- hwpx 2건은 hwp-mcp `create_hwpx_document`로 만든다(generate가 samples/_hwpx/*.json에 내용을 써 둔다).
- 스캔 PDF = 텍스트 PDF를 150dpi 회색 이미지로 구운 이미지 PDF. 카톡 JPG = 100dpi RGB, 폭 ≤1280px, JPEG 품질 70.
  원본 텍스트판은 samples/originals/ 에 남긴다.
"""
from __future__ import annotations

import json
import random
import re
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
INBOX = SAMPLES / "inbox"
ORIG = SAMPLES / "originals"
TEXT = SAMPLES / "text"
HWPX_JSON = SAMPLES / "_hwpx"
FONT = "/opt/hnc/hoffice11/Shared/TTF/Install/Hancom Gothic Regular.ttf"
FONT_FALLBACK = "/usr/share/fonts/noto-cjk/NotoSansCJK-Light.ttc"

random.seed(20260910)

DOC_TYPES = ["사업자등록증", "임대차계약서", "법인등기부등본", "통장내역", "급여대장", "근로계약서", "전자세금계산서",
             "종이세금계산서", "종이영수증", "4대보험취득확인서", "거래명세서", "카드매출매입내역", "잔액증명서",
             "원천세신고서", "부가세신고서"]
CATEGORY = {"사업자등록증": "기본서류", "임대차계약서": "기본서류", "법인등기부등본": "기본서류",
            "통장내역": "정기증빙", "급여대장": "정기증빙", "근로계약서": "정기증빙", "전자세금계산서": "정기증빙",
            "종이세금계산서": "정기증빙", "종이영수증": "정기증빙", "4대보험취득확인서": "정기증빙",
            "거래명세서": "신고증빙", "카드매출매입내역": "신고증빙", "잔액증명서": "신고증빙",
            "원천세신고서": "신고증빙", "부가세신고서": "신고증빙"}
SUB_MAP = {"급여대장": "01_원천세_급여", "근로계약서": "01_원천세_급여", "4대보험취득확인서": "01_원천세_급여", "원천세신고서": "01_원천세_급여",
           "통장내역": "02_통장_증빙", "전자세금계산서": "02_통장_증빙", "종이세금계산서": "02_통장_증빙", "종이영수증": "02_통장_증빙",
           "거래명세서": "03_부가세", "카드매출매입내역": "03_부가세", "부가세신고서": "03_부가세", "잔액증명서": "04_결산_조정"}

CLIENTS = {
    "hanbit": dict(name="한빛나루식당", short="한빛나루", kind="음식점", ceo="박정아", biz="999-81-00011", addr="광주광역시 북구 가상로 12", tel="062-000-0011", open="2021.03.15"),
    "daesung": dict(name="대성정밀공업사", short="대성정밀", kind="소규모 제조(금속가공)", ceo="이재훈", biz="999-81-00022", addr="광주광역시 광산구 가상산단길 34", tel="062-000-0022", open="2018.07.02"),
    "momo": dict(name="주식회사 모모스토어", short="모모스토어", kind="전자상거래(온라인몰)", ceo="정유진", biz="999-81-00033", addr="전남 나주시 가상혁신로 56", tel="061-000-0033", open="2022.11.21"),
    "pine": dict(name="푸른솔외국어학원", short="푸른솔", kind="교육서비스(학원)", ceo="최성민", biz="999-81-00044", addr="광주광역시 서구 가상학원길 78", tel="062-000-0044", open="2019.02.11"),
    "sejin": dict(name="세진건재상사", short="세진건재", kind="도소매(건설자재)", ceo="한동욱", biz="999-81-00055", addr="전남 담양군 가상건재로 90", tel="061-000-0055", open="2016.09.05"),
}
EMP = ["김민수", "이서연", "박지훈", "최유진", "정하늘", "오세진", "강다은"]

# (id, doc_type, client, period_normalized, cycle, format, filename, ocr_kind(None|scan|jpg), name_mode, note)
SPEC = [
    ("d01", "사업자등록증", "hanbit", "PERMANENT", "PERMANENT", "pdf", "사업자등록증.pdf", "scan", "full", "스캔본"),
    ("d02", "사업자등록증", "daesung", "PERMANENT", "PERMANENT", "jpg", "KakaoTalk_20260305_101522.jpg", "jpg", "full", "카톡 사진"),
    ("d03", "임대차계약서", "momo", "PERMANENT", "PERMANENT", "pdf", "스캔001.pdf", "scan", "full", "스캔본"),
    ("d04", "임대차계약서", "pine", "PERMANENT", "PERMANENT", "docx", "임대차계약서(수정본).docx", None, "full", ""),
    ("d05", "통장내역", "sejin", "2026-08", "MONTHLY", "xlsx", "거래내역조회_20260901.xlsx", None, "full", ""),
    ("d06", "통장내역", "hanbit", "2026-05", "MONTHLY", "xlsx", "가상은행_거래내역.xlsx", None, "short", "본문 약칭"),
    ("d07", "급여대장", "daesung", "2026-04", "MONTHLY", "xlsx", "급여 3월 최종(2).xlsx", None, "short", "파일명 3월·본문 4월 — 정답은 본문 · 약칭"),
    ("d08", "급여대장", "momo", "2026-07", "MONTHLY", "pdf", "7월 급여.pdf", None, "full", ""),
    ("d09", "근로계약서", "pine", "2026-03", "MONTHLY", "pdf", "스캔002.pdf", "scan", "full", "서명 스캔본"),
    ("d10", "근로계약서", "sejin", "2026-02", "MONTHLY", "docx", "근로계약서_정하늘.docx", None, "full", ""),
    ("d11", "종이세금계산서", "hanbit", "2026-02", "MONTHLY", "jpg", "KakaoTalk_20260212_183012.jpg", "jpg", "short", "카톡 사진 · 약칭"),
    ("d12", "종이세금계산서", "daesung", "2026-06", "MONTHLY", "pdf", "스캔003.pdf", "scan", "full", "스캔본"),
    ("d13", "거래명세서", "momo", "2026-Q2", "QUARTERLY", "pdf", "FAX_20260630.pdf", "scan", "full", "팩스 흉내(스캔)"),
    ("d14", "거래명세서", "pine", "2026-Q1", "QUARTERLY", "xlsx", "거래명세서 1분기.xlsx", None, "full", ""),
    ("d15", "카드매출매입내역", "sejin", "2026-Q2", "QUARTERLY", "xlsx", "카드매입매출_2분기.xlsx", None, "full", ""),
    ("d16", "카드매출매입내역", "hanbit", "2026-Q1", "QUARTERLY", "xlsx", "Book1.xlsx", None, "short", "본문 약칭"),
    ("d17", "잔액증명서", "daesung", "2025", "YEARLY", "pdf", "잔액증명.pdf", "scan", "full", "스캔본"),
    ("d18", "잔액증명서", "momo", "2025", "YEARLY", "pdf", "20251231_잔액증명서_모모.pdf", None, "full", "파일명에 약칭"),
    ("d19", "원천세신고서", "pine", "2026-05", "MONTHLY", "hwpx", "제출용.hwpx", None, "full", "사무소 산출물"),
    ("d20", "부가세신고서", "sejin", "2026-1H", "HALF", "pdf", "2026년 1기 부가세 신고서.pdf", None, "full", "사무소 산출물"),
    ("d21", "4대보험취득확인서", "hanbit", "2026-08", "MONTHLY", "hwpx", "취득확인.hwpx", None, "full", "사무소 산출물"),
    ("d22", "전자세금계산서", "momo", "2026-08", "MONTHLY", "pdf", "tax_invoice_0812.pdf", None, "full", ""),
]


def won(n: int) -> str:
    return f"{n:,}"


def period_parts(norm: str):
    """'2026-08' → (2026, 8) · '2026-Q2' → (2026, 6) · '2026-1H' → (2026, 6) · '2025' → (2025, 12) · PERMANENT → (2026, 1)"""
    if norm == "PERMANENT":
        return 2026, 1
    if "-" not in norm:
        return int(norm), 12
    y, v = norm.split("-")
    if v.startswith("Q"):
        return int(y), int(v[1]) * 3
    if v.endswith("H"):
        return int(y), 6 if v == "1H" else 12
    return int(y), int(v)


def target_path(client: str, category: str, doc_type: str, period_norm: str, ext: str) -> str:
    """리서치 §4.2 generate_target_path 규칙 그대로."""
    if category == "기본서류" or period_norm == "PERMANENT":
        return f"{client}/00_기본서류/{client}_{doc_type}.{ext}"
    year = period_norm.split("-")[0]
    sub = SUB_MAP.get(doc_type, "99_기타")
    return f"{client}/{year}/{sub}/{client}_{period_norm}_{doc_type}.{ext}"


# ---------------------------------------------------------------- 문서 내용
def build(doc_type: str, c: dict, norm: str, name_mode: str) -> dict:
    nm = c["name"] if name_mode == "full" else c["short"]
    r = random.Random(f"{doc_type}{c['name']}{norm}")
    y, m = period_parts(norm)
    ys = str(y)

    if doc_type == "사업자등록증":
        return dict(title="사 업 자 등 록 증  ( 일반과세자 )",
                    meta=[f"등록번호: {c['biz']}", f"상호(법인명): {nm}", f"대표자: {c['ceo']}", f"개업연월일: {c['open']}",
                          f"사업장 소재지: {c['addr']}", f"업태·종목: {c['kind']}", "발급사유: 신규", f"공동사업자: 없음"],
                    table=(["구분", "내용"], [["사업자 단위 과세 적용사업자 여부", "여( ) 부(V)"], ["전자세금계산서 전용 전자우편주소", "fake@example.invalid"]]),
                    tail=[f"{c['open'].split('.')[0]}년 {int(c['open'].split('.')[1])}월 {int(c['open'].split('.')[2])}일", "가상세무서장 (직인)"])
    if doc_type == "임대차계약서":
        dep = r.randint(10, 60) * 1000000
        rent = r.randint(80, 350) * 10000
        return dict(title="부동산(상가) 임대차계약서",
                    meta=[f"임대인(갑): 홍가상  (010-0000-0001)", f"임차인(을): {nm}  대표 {c['ceo']}  ({c['biz']})",
                          f"소재지: {c['addr']}  {r.randint(1,5)}층 {r.randint(101,110)}호  면적 {r.randint(40,180)}㎡", "1. 위 부동산을 아래 조건으로 임대차한다.",
                          f"2. 보증금 금 {won(dep)}원, 월 차임 금 {won(rent)}원(부가세 별도), 매월 {r.randint(1,28)}일 후불",
                          "3. 계약기간: 2025년 04월 01일부터 2027년 03월 31일까지 (24개월)"],
                    table=(["항목", "내용"], [["용도", c["kind"]], ["관리비", f"월 {won(r.randint(5,30)*10000)}원"], ["특약", "원상복구 의무, 전대 금지, 계약 만료 3개월 전 갱신 협의"]]),
                    tail=["본 계약을 증명하기 위하여 계약 당사자가 서명·날인 후 각 1통씩 보관한다.", "2025년 03월 20일", f"임대인 홍가상 (인)   임차인 {nm} {c['ceo']} (인)   중개: 가상공인중개사사무소"])
    if doc_type == "통장내역":
        rows = []
        bal = r.randint(5, 40) * 1000000
        for i in range(r.randint(14, 22)):
            d = r.randint(1, 28)
            out_, in_ = 0, 0
            if r.random() < 0.6:
                out_ = r.randint(3, 300) * 10000
            else:
                in_ = r.randint(10, 500) * 10000
            bal += in_ - out_
            rows.append([f"{ys}.{m:02d}.{d:02d} {r.randint(9,18):02d}:{r.randint(0,59):02d}", r.choice(["카드대금", "급여이체", "세금계산서입금", "임대료", "전기요금", "가상유통", "온라인매출정산", "4대보험", "부가세납부"]),
                         won(out_) if out_ else "", won(in_) if in_ else "", won(bal)])
        rows.sort()
        return dict(title="거래내역조회 (보통예금)",
                    meta=[f"예금주: {nm}", f"계좌번호: 000-0000-0000-{r.randint(10,99)} (가상은행)", f"조회기간: {ys}.{m:02d}.01 ~ {ys}.{m:02d}.{'30' if m in (4,6,9,11) else '31' if m != 2 else '28'}", "조회일시: " + f"{ys}.{m+1 if m<12 else 12:02d}.01 09:12", "단위: 원"],
                    table=(["거래일시", "적요", "출금액", "입금액", "잔액"], rows), tail=["※ 본 내역은 인터넷뱅킹 조회 결과이며 증명서가 아닙니다."])
    if doc_type == "급여대장":
        n = r.randint(4, 7)
        rows = []
        for i in range(n):
            base = r.randint(2100000, 3600000); ot = r.randint(0, 40) * 10000
            g = base + ot
            rows.append([EMP[i], r.choice(["사원", "주임", "대리", "과장"]), won(base), won(ot), won(g), won(int(g*0.045)), won(int(g*0.03545)), won(int(g*0.009)), won(int(g*0.028)), won(int(g*(1-0.045-0.03545-0.009-0.028)))])
        return dict(title=f"{ys}년 {m}월 급여대장 ({nm})",
                    meta=[f"사업장: {nm}", f"대표자: {c['ceo']}", f"귀속: {ys}년 {m:02d}월 (지급일 {ys}-{m:02d}-25)", f"인원: {n}명"],
                    table=(["성명", "직급", "기본급", "연장수당", "지급합계", "국민연금", "건강보험", "고용보험", "소득세", "실지급액"], rows),
                    tail=["※ 장기요양보험은 건강보험에 포함", "작성: 경리 담당    확인: 대표 (인)"])
    if doc_type == "근로계약서":
        d = r.randint(2, 20); emp = EMP[4]
        return dict(title="표 준 근 로 계 약 서 (기간의 정함이 없는 경우)",
                    meta=[f"{nm} (이하 \"사업주\"라 함)과(와) {emp} (이하 \"근로자\"라 함)은 다음과 같이 근로계약을 체결한다.",
                          f"1. 근로개시일(계약기간): {ys}년 {m:02d}월 {d:02d}일부터", f"2. 근무장소: {c['addr']}", f"3. 업무의 내용: {c['kind']} 관련 사무 및 보조",
                          "4. 소정근로시간: 09:00 ~ 18:00 (휴게 12:00 ~ 13:00)", "5. 근무일/휴일: 매주 월~금, 주휴일 매주 일요일"],
                    table=(["항목", "내용"], [["6. 임금", f"월급 {won(r.randint(2200000, 2900000))}원 (상여금 없음, 식대 200,000원)"], ["임금 지급일", "매월 25일 (휴일이면 전일)"], ["지급 방법", "근로자 명의 예금통장 입금"],
                                            ["7. 연차유급휴가", "근로기준법에 따라 부여"], ["8. 사회보험", "고용·산재·국민연금·건강보험 적용"], ["9. 계약서 교부", "체결과 동시에 사본 교부"]]),
                    tail=[f"{ys}년 {m:02d}월 {d:02d}일", f"(사업주) 사업체명: {nm}  대표자: {c['ceo']} (서명)  전화: {c['tel']}", f"(근로자) 성명: {emp} (서명)  연락처: 010-0000-0000"])
    if doc_type in ("전자세금계산서", "종이세금계산서"):
        amt = r.randint(12, 90) * 100000; d = r.randint(3, 27)
        supplier = r.choice(["가상유통 주식회사", "제일가상상사", "가상산업"])
        paper = doc_type == "종이세금계산서"
        return dict(title=("세 금 계 산 서 (공급받는자 보관용)" if paper else "전자세금계산서 (공급받는자 보관용)"),
                    meta=[f"작성일자: {ys}.{m:02d}.{d:02d}" if paper else f"작성일자: {ys}-{m:02d}-{d:02d}    승인번호: {ys}{m:02d}{d:02d}-41000000-0000{r.randint(1000,9999)}",
                          f"[공급자] 상호: {supplier}  등록번호: 999-86-{r.randint(10000,99999)}  성명: 홍가상",
                          f"[공급받는자] 상호: {nm}  등록번호: {c['biz']}  성명: {c['ceo']}", f"[공급받는자] 사업장 주소: {c['addr']}  업태·종목: {c['kind']}"],
                    table=(["월/일", "품목", "수량", "단가", "공급가액", "세액"], [[f"{m:02d}/{d:02d}", "식자재 일괄" if "음식" in c["kind"] else "원자재/소모품", "1", won(amt), won(amt), won(amt//10)], ["", "합계", "", "", won(amt), won(amt//10)]]),
                    tail=[f"합계금액: {won(amt+amt//10)}원 (공급가액 {won(amt)} + 세액 {won(amt//10)})", "이 금액을 (청구)함.", "수기 작성분 — 국세청 미전송" if paper else "국세청 전송 완료"])
    if doc_type == "거래명세서":
        q = (m - 1) // 3 + 1
        m0 = (q - 1) * 3 + 1
        rows = []
        items = ["절삭공구", "볼트/너트 세트", "포장 박스", "라벨지", "시멘트 40kg", "합판 12T", "교재 세트", "식자재"]
        for i in range(r.randint(6, 10)):
            qty = r.randint(1, 40); up = r.randint(3, 120) * 1000
            rows.append([f"{ys}.{r.randint(m0, m):02d}.{r.randint(1,28):02d}", r.choice(items), str(qty), won(up), won(qty*up), won(qty*up//10)])
        rows.sort(); s = sum(int(x[4].replace(",", "")) for x in rows)
        return dict(title="거 래 명 세 서",
                    meta=[f"공급받는자: {nm} 귀하  ({c['biz']})  대표: {c['ceo']}", f"거래기간: {ys}.{m0:02d}.01 ~ {ys}.{m:02d}.{'30' if m in (6,9) else '31'}  ({ys}년 {q}분기)",
                          "공급자: 가상종합상사 (999-86-00099)  대표: 홍가상", f"납품처: {c['addr']}"],
                    table=(["일자", "품명", "수량", "단가", "공급가액", "세액"], rows + [["", "합계", "", "", won(s), won(s//10)]]),
                    tail=[f"총 청구액: {won(s + s//10)}원", "결제조건: 익월 말 현금", f"발송일: {ys}.{m:02d}.30  인수자: (서명)"])
    if doc_type == "카드매출매입내역":
        q = (m - 1) // 3 + 1; m0 = (q - 1) * 3 + 1
        rows = []
        for i in range(r.randint(16, 24)):
            mm = r.randint(m0, m); d = r.randint(1, 28)
            kind = r.choice(["매출", "매출", "매입"])
            rows.append([f"{ys}-{mm:02d}-{d:02d}", kind, r.choice(["가상마트", "가상주유소", "손님결제", "손님결제", "가상문구", "가상통신"]), won(r.randint(8, 400) * 1000), "가상카드"])
        rows.sort()
        return dict(title=f"{c['short']}_가상카드_{q}분기",
                    meta=[f"가맹점·회원: {nm}  ({c['biz']})", f"조회기간: {ys}.{m0:02d}.01 ~ {ys}.{m:02d}.{'30' if m in (6,9) else '31'}", "구분: 매출(단말기 승인) / 매입(법인카드 사용)", "단위: 원 (VAT 포함)"],
                    table=(["승인일", "구분", "가맹점/거래처", "금액", "카드사"], rows), tail=["※ 부가세 신고용 카드 매출·매입 집계 자료"])
    if doc_type == "잔액증명서":
        bal = r.randint(12, 90) * 1000000
        return dict(title="예 금 잔 액 증 명 서",
                    meta=[f"예금주: {nm}", f"사업자등록번호: {c['biz']}", f"계좌번호: 000-0000-0000-{r.randint(10,99)} (보통예금)", f"발급기준일: {ys}년 12월 31일", f"증명 잔액: 금 {won(bal)}원 (₩{won(bal)})", "용도: 결산(재무제표 작성) 제출용"],
                    table=(["계좌", "종류", "잔액(원)", "질권·압류"], [[f"000-0000-0000-{r.randint(10,99)}", "보통예금", won(bal), "없음"]]),
                    tail=["위와 같이 잔액이 있음을 증명합니다.", f"{y+1}년 01월 {r.randint(5,20):02d}일", "가상은행 가상지점장 (직인)"])
    if doc_type == "원천세신고서":
        n = r.randint(3, 9); pay = n * r.randint(2200000, 3400000)
        return dict(title="[별지 제21호서식] 원천징수이행상황신고서",
                    meta=[f"귀속연월: {ys}/{m:02d}    지급연월: {ys}/{m:02d}", "매월 신고 · 원천징수의무자 제출", f"법인명(상호): {nm}", f"대표자: {c['ceo']}", f"사업자등록번호: {c['biz']}", f"소재지: {c['addr']}", f"전화: {c['tel']}"],
                    table=(["소득종류", "코드", "인원", "총지급액(원)", "소득세(원)", "농어촌특별세"], [["근로소득 간이세액", "A01", str(n), won(pay), won(int(pay*0.028)), "0"], ["근로소득 중도퇴사", "A02", "0", "0", "0", "0"], ["사업소득", "A25", str(r.randint(0,2)), won(r.randint(0,3)*800000), won(r.randint(0,3)*24000), "0"], ["합계", "", str(n), "", won(int(pay*0.028)), "0"]]),
                    tail=[f"신고일: {ys}년 {m+1 if m<12 else 1:02d}월 {r.randint(5,10)}일", f"신고인: {c['ceo']} (인)", "세무대리인: 가상세무회계사무소 (관리번호 000-000)"])
    if doc_type == "부가세신고서":
        half = "1기 확정" if m <= 6 else "2기 확정"
        sales = r.randint(80, 400) * 100000; purch = int(sales * r.uniform(0.45, 0.7)) // 10000 * 10000
        return dict(title=f"[별지 제21호서식] 일반과세자 부가가치세 {half} 신고서",
                    meta=[f"과세기간: {ys}년 {half} ({'01.01~06.30' if m<=6 else '07.01~12.31'})", f"상호(법인명): {nm}", f"성명(대표자): {c['ceo']}", f"사업자등록번호: {c['biz']}", f"사업장 소재지: {c['addr']}", f"전화번호: {c['tel']}", f"업태·종목: {c['kind']}"],
                    table=(["구분", "금액(원)", "세율", "세액(원)"], [["과세 세금계산서 발급분", won(sales), "10/100", won(sales//10)], ["과세 기타(카드·현금영수증)", won(sales//3), "10/100", won(sales//30)], ["매입 세금계산서 수취분", won(purch), "", won(purch//10)], ["매입 기타 공제", won(purch//8), "", won(purch//80)], ["납부(환급)세액", "", "", won(sales//10 + sales//30 - purch//10 - purch//80)]]),
                    tail=["신고인은 「부가가치세법」 제48조·제49조에 따라 위 내용을 신고하며 사실 그대로 정확하게 적었음을 확인합니다.", f"신고일: {ys}년 {7 if m<=6 else 1}월 {r.randint(10,24)}일    신고인: {c['ceo']} (서명 또는 인)", "세무대리인: 가상세무회계사무소"])
    if doc_type == "4대보험취득확인서":
        d = r.randint(2, 20); emp = EMP[2]
        return dict(title="사업장가입자 자격취득 확인서 (국민연금·건강보험·고용보험·산재보험)",
                    meta=[f"사업장명칭: {nm}", f"사업장관리번호: {c['biz']}-0", f"사업자등록번호: {c['biz']}", f"소재지: {c['addr']}", f"전화: {c['tel']}", f"처리일: {ys}년 {m:02d}월 {d+3:02d}일"],
                    table=(["성명", "주민등록번호", "자격취득일", "월소득액(원)", "국민연금", "건강보험", "고용보험", "산재보험"], [[emp, "000000-0000000", f"{ys}.{m:02d}.{d:02d}", won(r.randint(2200000, 3000000)), "취득", "취득", "취득", "취득"]]),
                    tail=["위와 같이 자격취득이 처리되었음을 확인합니다.", "국민연금공단·국민건강보험공단·근로복지공단 (4대사회보험 정보연계센터)", "※ 사무소 보관용 출력본"])
    raise ValueError(doc_type)


# ---------------------------------------------------------------- 렌더러
def _font():
    return FONT if Path(FONT).exists() else FONT_FALLBACK


def render_pdf(doc: dict, out: Path) -> None:
    import pymupdf
    pdf = pymupdf.open(); page = pdf.new_page(width=595, height=842)
    font, fn = _font(), "kfont"
    y = 60
    page.insert_text((50, y), doc["title"], fontsize=14, fontname=fn, fontfile=font); y += 28
    for line in doc["meta"]:
        for chunk in wrap(line, 62):
            page.insert_text((50, y), chunk, fontsize=9.5, fontname=fn, fontfile=font); y += 15
    y += 8
    headers, rows = doc["table"]
    x0, x1 = 50, 545
    widths = col_widths(headers, rows, x1 - x0); rh = 16
    xs = [x0]
    for w in widths:
        xs.append(xs[-1] + w)
    def draw_row(vals, yy):
        page.draw_line((x0, yy), (x1, yy), width=0.5)
        for i, v in enumerate(vals):
            page.insert_text((xs[i] + 3, yy + 12), str(v)[:22], fontsize=8, fontname=fn, fontfile=font)
    draw_row(headers, y); y += rh
    for rrow in rows:
        if y > 780:
            break
        draw_row(rrow, y); y += rh
    page.draw_line((x0, y), (x1, y), width=0.5)
    for xx in xs:
        page.draw_line((xx, y - rh * (min(len(rows), (780 - 60) // rh) + 1)), (xx, y), width=0.5)
    y += 20
    for line in doc["tail"]:
        for chunk in wrap(line, 62):
            page.insert_text((50, y), chunk, fontsize=9.5, fontname=fn, fontfile=font); y += 15
    pdf.subset_fonts()
    pdf.save(out, garbage=4, deflate=True)


def render_scanned(src_pdf: Path, out: Path, dpi: int = 150) -> None:
    import pymupdf
    src = pymupdf.open(src_pdf); dst = pymupdf.open()
    for p in src:
        pix = p.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        page = dst.new_page(width=p.rect.width, height=p.rect.height)
        page.insert_image(page.rect, pixmap=pix)
    dst.save(out)


def render_kakao_jpg(src_pdf: Path, out: Path) -> None:
    """카톡으로 받은 사진 흉내: 100dpi RGB → 폭 ≤1280px → JPEG 품질 70."""
    import pymupdf
    src = pymupdf.open(src_pdf)
    p = src[0]
    pix = p.get_pixmap(dpi=100, colorspace=pymupdf.csRGB)
    if pix.width > 1280:
        pix = p.get_pixmap(matrix=pymupdf.Matrix(1280 / p.rect.width, 1280 / p.rect.width), colorspace=pymupdf.csRGB)
    pix.save(out, jpg_quality=70)


def render_xlsx(doc: dict, out: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = doc["title"][:28].replace("/", "_")
    ws.append([doc["title"]])
    for line in doc["meta"]:
        ws.append([line])
    ws.append([])
    headers, rows = doc["table"]
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.append([])
    for line in doc["tail"]:
        ws.append([line])
    wb.save(out)


def render_docx(doc: dict, out: Path) -> None:
    import docx
    d = docx.Document(); d.add_heading(doc["title"], level=1)
    for line in doc["meta"]:
        d.add_paragraph(line)
    headers, rows = doc["table"]
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for line in doc["tail"]:
        d.add_paragraph(line)
    d.save(out)


def hwpx_content(doc: dict) -> list:
    items = [{"type": "text", "text": doc["title"]}] + [{"type": "text", "text": l} for l in doc["meta"]]
    headers, rows = doc["table"]
    items.append({"type": "table", "headers": headers, "rows": [[str(v) for v in r] for r in rows]})
    return items + [{"type": "text", "text": l} for l in doc["tail"]]


def wrap(s: str, n: int) -> list[str]:
    return [s[i:i + n] for i in range(0, len(s), n)] or [""]


def col_widths(headers, rows, total):
    n = len(headers)
    lens = [max([len(str(headers[i]))] + [len(str(r[i])) for r in rows]) for i in range(n)]
    lens = [max(3, min(l, 22)) for l in lens]
    s = sum(lens)
    return [total * l / s for l in lens]


# ---------------------------------------------------------------- 추출
# 파서는 엔진이 정본이다 — 여기서 다시 구현하지 않는다(2026-09-10, hwpx 줄바꿈 결함이 두 곳에 있었다).
def extract_text(path: Path) -> str:
    sys.path.insert(0, str(ROOT))
    from engine.extract import extract as _extract
    return _extract(path).text


# ---------------------------------------------------------------- main
def generate() -> None:
    for d in (INBOX, ORIG, TEXT, HWPX_JSON):
        d.mkdir(parents=True, exist_ok=True)
    for old in list(INBOX.iterdir()) + list(ORIG.iterdir()) + list(HWPX_JSON.iterdir()):
        if old.is_file() and old.suffix.lower() != ".hwpx":
            old.unlink()
    for (did, dt, ck, norm, cycle, fmt, fname, ocr, nmode, note) in SPEC:
        c = CLIENTS[ck]; doc = build(dt, c, norm, nmode); out = INBOX / fname
        if fmt == "pdf" and ocr == "scan":
            orig = ORIG / f"{did}-{fname}"; render_pdf(doc, orig); render_scanned(orig, out)
        elif fmt == "jpg":
            orig = ORIG / f"{did}-{Path(fname).stem}.pdf"; render_pdf(doc, orig); render_kakao_jpg(orig, out)
        elif fmt == "pdf":
            render_pdf(doc, out)
        elif fmt == "xlsx":
            render_xlsx(doc, out)
        elif fmt == "docx":
            render_docx(doc, out)
        elif fmt == "hwpx":
            (HWPX_JSON / f"{did}.json").write_text(json.dumps({"output_path": str(out), "content": hwpx_content(doc)}, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{did} {fmt:4} {ocr or '    ':4} {fname}")
    print("hwpx 2건은 samples/_hwpx/*.json 을 hwp-mcp create_hwpx_document 에 넘겨 만든다.")


def extract() -> None:
    TEXT.mkdir(parents=True, exist_ok=True)
    labels = []
    for (did, dt, ck, norm, cycle, fmt, fname, ocr, nmode, note) in SPEC:
        c = CLIENTS[ck]; p = INBOX / fname
        if not p.exists():
            print(f"!! 없음: {fname}"); continue
        txt = extract_text(p)
        (TEXT / f"{did}.txt").write_text(txt, encoding="utf-8")
        cat = CATEGORY[dt]
        labels.append(dict(id=did, file=fname, client=c["name"], client_short=c["short"], category=cat, doc_type=dt,
                           period=dict(cycle=cycle, normalized=norm), format=fmt, scanned=bool(ocr), ocr_kind=ocr, note=note,
                           target_path=target_path(c["name"], cat, dt, norm, fmt), chars=len(txt)))
        print(f"{did} {len(txt):5d}자 {'OCR' if ocr else '   '} {fname}")
    (SAMPLES / "labels.json").write_text(json.dumps(labels, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"labels.json {len(labels)}건 · OCR {sum(1 for l in labels if l['scanned'])}건")


if __name__ == "__main__":
    {"generate": generate, "extract": extract}[sys.argv[1] if len(sys.argv) > 1 else "generate"]()
